"""Обработка выгрузки КИЗ из личного кабинета Wildberries.

Лист "КИЗ" содержит служебные столбцы выгрузки; нужны только КИЗ и
стоимость. Значение КИЗ WB "склеивает" из нескольких полей DataMatrix-кода
(группы AI 01/21, 91, 92) через служебный разделитель, который при экспорте
в xlsx превращается в символ U+FFFD (отрисовывается как чёрный ромб) — нам
нужна только первая группа, до этого разделителя.
"""

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.workbook import Workbook as WorkbookType

SHEET_NAME = "КИЗ"
KIZ_COLUMN = "КИЗ"
PRICE_COLUMN = "Стоимость"
KIZ_SEPARATOR = "�"


def detect(workbook: WorkbookType) -> bool:
    if SHEET_NAME not in workbook.sheetnames:
        return False
    headers = [cell.value for cell in workbook[SHEET_NAME][1]]
    return KIZ_COLUMN in headers and PRICE_COLUMN in headers


def process(workbook: WorkbookType) -> WorkbookType:
    sheet = workbook[SHEET_NAME]
    headers = [cell.value for cell in sheet[1]]
    kiz_idx = headers.index(KIZ_COLUMN)
    price_idx = headers.index(PRICE_COLUMN)

    result = Workbook()
    result_sheet = result.active
    result_sheet.title = SHEET_NAME
    result_sheet.append([KIZ_COLUMN, PRICE_COLUMN])
    for cell in result_sheet[1]:
        cell.font = Font(bold=True)

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[kiz_idx] is None:
            continue
        kiz = str(row[kiz_idx]).split(KIZ_SEPARATOR, 1)[0]
        result_sheet.append([kiz, row[price_idx]])

    for column_cells in result_sheet.columns:
        width = max(
            len(str(cell.value)) for cell in column_cells if cell.value is not None
        )
        result_sheet.column_dimensions[column_cells[0].column_letter].width = width + 2

    return result
